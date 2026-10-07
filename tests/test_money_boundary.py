"""NEX-54 — money crosses the API boundary as an exact decimal string.

One test per endpoint that used to cast a Decimal to float on the way out.
(The dashboard has its own leaf-by-leaf walk in test_dashboard_concurrent.py.)
"""
import csv
import io
import json
import re
from datetime import datetime, timedelta, timezone
from decimal import Decimal as D

import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient

from app.core.gold_api import GoldRateResult, build_rate_history_row
from app.models import (
    CoinType, GoldRateHistory, GoldRateOverride, Karat, MarginMode, Order, OrderStatus, OunceType,
    PaymentMethod, Product, ProductStatus, Role, Settings, User,
)


@pytest_asyncio.fixture
async def client(db):
    from app.deps import get_current_user, get_db
    from app.main import app

    admin = User(id="u-admin", email="a@x.com", name="Admin", password_hash="x",
                 role=Role.ADMIN, is_active=True)
    db.add_all([admin, Settings(id="singleton")])
    await db.flush()

    async def _get_db():
        yield db

    async def _get_user():
        return admin

    app.dependency_overrides[get_db] = _get_db
    app.dependency_overrides[get_current_user] = _get_user
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        yield c
    app.dependency_overrides.clear()


async def _rate(db, value: str, **cols):
    db.add(GoldRateHistory(rate_24k=D(value), source="goldapi", **cols))
    await db.flush()


# ── gold price ────────────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_current_gold_rate_is_strings_for_every_karat(client, db):
    await _rate(db, "84.31")                # no per-karat values on the row: derived
    body = (await client.get("/api/gold-price")).json()
    assert body["rate_24k"] == "84.31"
    assert body["rate_22k"] == "77.31"      # 84.31 × 0.917 = 77.31227
    assert body["rate_21k"] == "73.77"      # 84.31 × 0.875 = 73.77125
    assert body["rate_18k"] == "63.23"      # 84.31 × 0.750 = 63.2325
    assert body["is_stale"] is False and body["source"] == "live"

    await _rate(db, "80.10", fetched_at=datetime.now(timezone.utc) + timedelta(seconds=1))
    assert (await client.get("/api/gold-price")).json()["rate_24k"] == "80.10"   # trailing zero kept


@pytest.mark.asyncio
async def test_live_gold_rate_returns_the_per_karat_values_stored_on_the_row(client, db):
    """The poller derives the karats from the unrounded feed price, so what it
    stores need not equal 24K × purity. The card must show what is stored — the
    same figures as the latest point on the chart — not a second derivation."""
    polled = build_rate_history_row(D("84.3149"), "goldapi")
    assert (polled.rate_22k, polled.rate_21k, polled.rate_18k) == (D("77.32"), D("73.78"), D("63.24"))
    # The row as the database holds it: NUMERIC(10,2) keeps 84.31 of the 24K price.
    db.add(GoldRateHistory(rate_24k=D("84.31"), rate_22k=polled.rate_22k, rate_21k=polled.rate_21k,
                           rate_18k=polled.rate_18k, source="goldapi",
                           fetched_at=datetime.now(timezone.utc) - timedelta(seconds=5)))
    await db.flush()

    card = (await client.get("/api/gold-price")).json()
    latest_point = (await client.get("/api/gold-price/history?range=24h")).json()[-1]

    assert card["rate_24k"] == latest_point["rate_24k"] == "84.31"
    # Derived from the stored 84.31 these would be 77.31 / 73.77 / 63.23.
    assert (card["rate_22k"], card["rate_21k"], card["rate_18k"]) == ("77.32", "73.78", "63.24")
    assert all(card[k] == latest_point[k] for k in ("rate_22k", "rate_21k", "rate_18k"))


@pytest.mark.asyncio
async def test_derived_karats_use_the_pricing_engines_purity_and_rounding(client, db):
    """An override has no stored karats. Deriving them must give what the
    pricing engine quotes: 84.30 × 0.750 = 63.225 is 63.23, not 63.22."""
    await _rate(db, "80.00")
    db.add(GoldRateOverride(rate_24k=D("84.30"), set_by="u-admin", is_active=True))
    db.add(Product(code="RNG-18", name_en="Ring", category="rings", karat=Karat.K18,
                   weight_grams=D("5"), margin_percent=D("15"), making_charge=D("20"),
                   on_hand_qty=1, is_active=True, status=ProductStatus.AVAILABLE))
    await db.flush()

    card = (await client.get("/api/gold-price")).json()
    assert card["source"] == "override"
    assert (card["rate_24k"], card["rate_22k"], card["rate_21k"], card["rate_18k"]) == (
        "84.30", "77.30", "73.76", "63.23")
    lookup = (await client.get("/api/products/lookup/RNG-18")).json()
    assert lookup["purity_rate"] == card["rate_18k"] == "63.23"


@pytest.mark.asyncio
async def test_derived_karats_read_the_shared_purity_table(client, db, monkeypatch):
    from app.core import pricing

    await _rate(db, "100.00")
    monkeypatch.setitem(pricing.KARAT_PURITY, Karat.K22, D("0.500"))     # no private copy to miss this
    assert (await client.get("/api/gold-price")).json()["rate_22k"] == "50.00"
    history = (await client.get("/api/gold-price/history?range=24h")).json()
    assert history[-1]["rate_22k"] == "50.00"


@pytest.mark.asyncio
async def test_active_override_rate_is_a_string(client, db):
    await _rate(db, "84.31")
    db.add(GoldRateOverride(rate_24k=D("90"), set_by="u-admin", is_active=True))
    await db.flush()
    body = (await client.get("/api/gold-price")).json()
    assert body["source"] == "override"
    assert (body["rate_24k"], body["rate_18k"]) == ("90.00", "67.50")


@pytest.mark.asyncio
async def test_gold_rate_history_points_are_strings(client, db):
    row = build_rate_history_row(D("80"), "goldapi")
    row.fetched_at = datetime.now(timezone.utc) - timedelta(hours=1)
    # A row from before per-karat storage: the karats are derived on the way out.
    legacy = GoldRateHistory(rate_24k=D("100"), source="goldapi",
                             fetched_at=datetime.now(timezone.utc) - timedelta(hours=2))
    db.add_all([row, legacy])
    await db.flush()

    legacy_point, point = (await client.get("/api/gold-price/history?range=24h")).json()
    assert {k: point[k] for k in ("rate_24k", "rate_22k", "rate_21k", "rate_18k")} == {
        "rate_24k": "80.00", "rate_22k": "73.36", "rate_21k": "70.00", "rate_18k": "60.00"}
    assert {k: legacy_point[k] for k in ("rate_24k", "rate_22k", "rate_21k", "rate_18k")} == {
        "rate_24k": "100.00", "rate_22k": "91.70", "rate_21k": "87.50", "rate_18k": "75.00"}


@pytest.mark.asyncio
async def test_refresh_and_override_echo_the_rate_as_a_string(client, db, monkeypatch):
    from app.api import gold_price

    async def _fetch():
        return GoldRateResult(value=D("84.3127"), source="goldapi")

    monkeypatch.setattr(gold_price, "fetch_gold_rate", _fetch)
    r = await client.post("/api/gold-price/refresh")
    assert r.status_code == 200, r.text
    assert r.json() == {"rate": "84.31", "source": "goldapi"}

    r = await client.post("/api/gold-price/override", json={"rate_24k": "85.5", "reason": "feed down"})
    assert r.status_code == 200, r.text
    assert r.json() == {"message": "Override set", "rate_24k": "85.50"}

    # Half-cent ties round up, which is what Postgres does when it stores the
    # NUMERIC(10,2): the echo and the stored rate cannot be a cent apart.
    async def _fetch_tie():
        return GoldRateResult(value=D("84.325"), source="goldapi")

    monkeypatch.setattr(gold_price, "fetch_gold_rate", _fetch_tie)
    assert (await client.post("/api/gold-price/refresh")).json()["rate"] == "84.33"
    r = await client.post("/api/gold-price/override", json={"rate_24k": "85.505", "reason": "feed down"})
    assert r.json()["rate_24k"] == "85.51"


# ── price previews ────────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_price_previews_quote_the_gold_rate_as_a_string(client, db):
    await _rate(db, "84.30")
    coin = CoinType(code="LIRA-8", name_en="Lira", karat=Karat.K21, weight_grams=D("8"),
                    margin_mode=MarginMode.USD, on_hand_qty=3)
    bar = OunceType(code="OZ-1", name_en="Ounce", karat=Karat.K24, weight_grams=D("31.104"),
                    margin_mode=MarginMode.USD, on_hand_qty=3)
    ring = Product(code="RNG-1", name_en="Ring", category="rings", karat=Karat.K18,
                   weight_grams=D("5"), margin_percent=D("15"), making_charge=D("20"),
                   on_hand_qty=1, is_active=True, status=ProductStatus.AVAILABLE)
    db.add_all([coin, bar, ring])
    await db.flush()

    for url in (f"/api/coins/{coin.id}/price", f"/api/ounces/{bar.id}/price",
                "/api/products/lookup/RNG-1"):
        r = await client.get(url)
        assert r.status_code == 200, r.text
        assert r.json()["gold_rate_24k"] == "84.30", url
        assert b'"gold_rate_24k":"84.30"' in r.content


# ── stale-rate guard ──────────────────────────────────────────────────────────

def test_stale_rate_409_carries_the_rate_as_a_string():
    from fastapi import HTTPException

    from app.core.gold_guard import assert_rate_acceptable

    rate_info = {"rate": 84.3, "source": "live", "is_stale": True, "market_closed": True,
                 "fetched_at": datetime.now(timezone.utc) - timedelta(hours=3)}
    with pytest.raises(HTTPException) as exc:
        assert_rate_acceptable(rate_info, None)
    assert exc.value.detail["rate_24k"] == "84.30"
    assert '"rate_24k": "84.30"' in json.dumps(exc.value.detail)


# ── orders ────────────────────────────────────────────────────────────────────

def _order(number, total, *, lbp=None, status=OrderStatus.COMPLETED, minutes_ago=1):
    return Order(order_number=number, cashier_id="u-admin", status=status,
                 payment_method=PaymentMethod.CASH, subtotal=total, vat_percent=D("0"),
                 vat_amount=D("0"), total_usd=total, total_lbp=lbp or total * D("89500"),
                 lbp_exchange_rate=D("89500"),
                 created_at=datetime.now(timezone.utc) - timedelta(minutes=minutes_ago))


@pytest.mark.asyncio
async def test_order_list_totals_are_two_decimal_strings(client, db):
    db.add_all([_order("O-1", D("100.10"), minutes_ago=3), _order("O-2", D("0.10"), minutes_ago=2),
                _order("O-3", D("50.00"), minutes_ago=1)])
    await db.flush()
    body = (await client.get("/api/orders")).json()
    # Format only: the amounts come from the revenue query, which NEX-48 is
    # fixing on its own branch. A quotient used to go out at 28 digits.
    assert re.fullmatch(r"\d+\.\d{2}", body["total_revenue"]), body["total_revenue"]
    assert re.fullmatch(r"\d+\.\d{2}", body["avg_order_value"]), body["avg_order_value"]
    assert [o["total_usd"] for o in body["items"]] == ["50.00", "0.10", "100.10"]


def test_order_average_rounds_a_half_cent_up():
    # On the schema, so it does not depend on the revenue query (NEX-48):
    # 0.25 over two orders is 0.125.
    from app.schemas.order import OrderListOut

    out = OrderListOut(items=[], total=2, total_revenue=D("0.25"), avg_order_value=D("0.25") / 2)
    assert json.loads(out.model_dump_json()) == {
        "items": [], "total": 2, "total_revenue": "0.25", "avg_order_value": "0.13"}


@pytest.mark.asyncio
async def test_empty_order_list_totals_are_zero_strings(client):
    body = (await client.get("/api/orders")).json()
    assert (body["total_revenue"], body["avg_order_value"]) == ("0.00", "0.00")


@pytest.mark.asyncio
async def test_orders_csv_export_writes_exact_amounts(client, db):
    # (SQLite keeps NUMERIC as a double, so the fixture stays within 15 digits;
    # test_money.py covers a full NUMERIC(18,2) value, which a float corrupts.)
    db.add(_order("O-9", D("1234.50"), lbp=D("110487750000.50")))
    await db.flush()
    r = await client.get("/api/orders/export")
    assert r.status_code == 200
    header, row = list(csv.reader(io.StringIO(r.text)))
    cells = dict(zip(header, row))
    assert cells["Subtotal USD"] == "1234.50"
    assert cells["VAT"] == "0.00"
    assert cells["Total USD"] == "1234.50"
    assert cells["Total LBP"] == "110487750000.50"


# ── Excel exports are the exception: numbers stay numbers ─────────────────────

def test_xlsx_exports_still_write_numeric_cells():
    """A spreadsheet cell is typed, unlike JSON: money written as text would
    stop summing in Excel. xlsx.py keeps converting Decimal to a number."""
    import openpyxl

    from app.core.xlsx import Sheet, build_xlsx_bytes

    body = build_xlsx_bytes([Sheet(name="Money", headers=["Account", "Debit (USD)"],
                                   rows=[["Cash", D("1234.50")], ["Bank", D("0.10")], ["Total", D("1234.60")]])])
    ws = openpyxl.load_workbook(io.BytesIO(body))["Money"]
    amounts = [ws.cell(row=r, column=2) for r in (2, 3, 4)]
    assert [c.value for c in amounts] == [1234.5, 0.1, 1234.6]
    assert all(c.data_type == "n" for c in amounts)              # numeric, not "s" (string)
    assert ws.cell(row=2, column=1).data_type == "s"
